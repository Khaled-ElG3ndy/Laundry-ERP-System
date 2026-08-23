# -*- coding: utf-8 -*-
"""Regenerate ``data/price_book.py`` from the hotel price sheet.

Run this whenever the Excel file changes, then upgrade the module::

    python3 custom_addons/hosny_hotel_pos/tools/import_price_book.py \
        "/opt/odoo17/كشف اسعار مفصل بالعماء 1.xlsx"
    sudo systemctl stop odoo17-farha
    sudo -u odoo17 /opt/odoo17/venv/bin/python /opt/odoo17/odoo/odoo-bin \
        -c /etc/odoo17-farha.conf -u hosny_hotel_pos --stop-after-init
    sudo systemctl start odoo17-farha

The sheet ``كشف فواتير يومي`` holds four price blocks side by side.  Each block
is (index, item name, price) with its own row offset; the customers a block
serves are written in the merged header cells above it.  Item identity across
blocks comes from the normalised Arabic name, not the row index -- the blocks
do not share a row offset.

Only prices are regenerated.  Product identity (``code``) and the English names
are keyed off the normalised Arabic name and preserved across runs, so existing
products are updated in place rather than duplicated.
"""

import re
import sys
import unicodedata
from datetime import datetime, timezone

try:
    import openpyxl
except ImportError:  # pragma: no cover
    sys.exit("openpyxl is required: pip install openpyxl")

SHEET = "كشف فواتير يومي"

# tier -> (index col, name col, price col, first row, last row)
BLOCKS = {
    "base": (1, 2, 3, 6, 41),
    "fann": (16, 17, 18, 7, 42),
    "naseem": (21, 22, 23, 7, 42),
    "lulua": (26, 27, 28, 6, 41),
}

# Where each block writes the customers it serves.
CUSTOMER_CELLS = {
    "base": [(4, col) for col in range(4, 15)],  # D4 .. N4
    "fann": [(5, 19)],                           # S5
    "naseem": [(5, 24)],                         # X5
    "lulua": [(4, 29)],                          # AC4
}

# Normalised Arabic name -> (product code, English name).  The codes are what
# make the import repeatable: they become ``default_code`` on the product, so a
# second run updates instead of duplicating.  Add a row here when the sheet
# gains an item; never renumber an existing one.
CATALOG = [
    ("شرشف سرير كبير", "HTL-01", "Large Bed Sheet"),
    ("شرشف سرير صغير", "HTL-02", "Small Bed Sheet"),
    ("غطاء سرير كبير", "HTL-03", "Large Duvet Cover"),
    ("غطاء سرير صغير", "HTL-04", "Small Duvet Cover"),
    ("غطاء سرير عازل للماء كبير", "HTL-05", "Large Waterproof Mattress Protector"),
    ("غطاء سرير عازل للماء صغير", "HTL-06", "Small Waterproof Mattress Protector"),
    ("لحاف كبير", "HTL-07", "Large Duvet"),
    ("لحاف صغير", "HTL-08", "Small Duvet"),
    ("بطانية", "HTL-09", "Blanket"),
    ("مخدة", "HTL-10", "Pillow"),
    ("كيس مخدة", "HTL-11", "Pillowcase"),
    ("منشفة كبير (جسم)", "HTL-12", "Large Towel (Body)"),
    ("منشفة وسط (يد)", "HTL-13", "Medium Towel (Hand)"),
    ("منشفة صغير(وجه)", "HTL-14", "Small Towel (Face)"),
    ("دواسة", "HTL-15", "Bath Mat"),
    ("منشفة روب", "HTL-16", "Bathrobe"),
    ("حلية للسرير كبير", "HTL-17", "Large Bed Valance"),
    ("حلية للسرير صغير", "HTL-18", "Small Bed Valance"),
    ("غطاء كرسي", "HTL-19", "Chair Cover"),
    ("مفرش طاولة كبير", "HTL-20", "Large Tablecloth"),
    ("مفرش طاولة صغير", "HTL-21", "Small Tablecloth"),
    ("منديل طعام", "HTL-22", "Napkin"),
    ("ديكور طاولة", "HTL-23", "Table Runner"),
    ("حلية طاولة", "HTL-24", "Table Skirt"),
    ("قميص", "HTL-25", "Shirt"),
    ("بنطلون", "HTL-26", "Trousers"),
    ("صديري", "HTL-27", "Vest"),
    ("تي شيرت", "HTL-28", "T-Shirt"),
    ("جاكيت", "HTL-29", "Jacket"),
    ("جاكيت مطبخ", "HTL-30", "Chef Jacket"),
    ("فستان", "HTL-31", "Dress"),
    ("ستارة شيفون", "HTL-32", "Chiffon Curtain (per m)"),
    ("ستارة عازل", "HTL-33", "Blackout Curtain (per m)"),
    ("ستارة قماش", "HTL-34", "Fabric Curtain (per m)"),
    ("سجادة صلاة", "HTL-35", "Prayer Rug"),
    ("بدلة كاملة ( ثلاث قطع )", "HTL-36", "Full Suit (3 Pieces)"),
]

_DIACRITICS = re.compile(r"[ؐ-ًؚ-ٰٟۖ-ۭـ]")
_NON_WORD = re.compile(r"[^\w؀-ۿ]+", re.UNICODE)
_EQUIVALENTS = str.maketrans({
    "أ": "ا", "إ": "ا", "آ": "ا", "ى": "ي", "ة": "ه", "ؤ": "و", "ئ": "ي",
})


def normalise(value):
    """Fold the Arabic spelling variants the sheet mixes freely.

    The same item is written ``تي شيرت `` in one block and ``تي شيرت`` in
    another, and hamza/ta-marbuta forms drift between blocks.  Folding those
    plus whitespace is enough to key items across blocks; nothing fuzzier is
    used, so a genuinely new item shows up as unmatched rather than silently
    absorbing another item's price.
    """
    if value is None:
        return ""
    text = unicodedata.normalize("NFKC", str(value))
    text = _DIACRITICS.sub("", text).translate(_EQUIVALENTS)
    return " ".join(_NON_WORD.sub(" ", text).split()).lower()


def parse(path):
    workbook = openpyxl.load_workbook(path, data_only=True)
    if SHEET not in workbook.sheetnames:
        sys.exit("sheet %r not found in %s" % (SHEET, path))
    sheet = workbook[SHEET]

    customers = []
    for tier, cells in CUSTOMER_CELLS.items():
        for row, col in cells:
            raw = sheet.cell(row=row, column=col).value
            if raw and str(raw).strip():
                customers.append((tier, " ".join(str(raw).split())))

    prices = {}
    for tier, (_idx_col, name_col, price_col, first, last) in BLOCKS.items():
        for row in range(first, last + 1):
            name = sheet.cell(row=row, column=name_col).value
            if not name or not str(name).strip() or str(name).strip() == "0":
                continue
            key = normalise(name)
            if not key:
                continue
            price = sheet.cell(row=row, column=price_col).value
            if isinstance(price, (int, float)):
                prices.setdefault(key, {})[tier] = round(float(price), 2)
    return customers, prices


def build(path):
    customers, prices = parse(path)
    catalog = {normalise(ar): (ar, code, en) for ar, code, en in CATALOG}

    unknown = sorted(set(prices) - set(catalog))
    missing = sorted(set(catalog) - set(prices))
    if unknown:
        sys.stderr.write(
            "WARNING: %d item(s) in the sheet are absent from CATALOG and were "
            "skipped -- add them with a fresh HTL-xx code:\n  %s\n"
            % (len(unknown), "\n  ".join(unknown))
        )
    if missing:
        sys.stderr.write(
            "WARNING: %d CATALOG item(s) are absent from the sheet:\n  %s\n"
            % (len(missing), "\n  ".join(missing))
        )

    rows = []
    for ar, code, en in CATALOG:
        tier_prices = prices.get(normalise(ar))
        if tier_prices is None:
            continue
        rows.append((code, ar, en, tier_prices))
    return customers, rows


TEMPLATE = '''# -*- coding: utf-8 -*-
"""Generated by tools/import_price_book.py -- do not edit by hand.

Source: {source}
Generated: {stamp}

PRODUCTS is (code, arabic_name, english_name, {{tier: price}}).  ``code``
becomes ``default_code`` on the product and is the key the sync matches on, so
re-running the import updates rather than duplicates.

SHEET_CUSTOMERS is the customer list exactly as the sheet spells it, kept for
traceability; the mapping onto real partners lives in models/hotel_setup.py.
"""

TIERS = {tiers!r}

PRODUCTS = [
{products}]

SHEET_CUSTOMERS = [
{customers}]
'''


def render(source, customers, rows):
    products = "".join(
        "    (%r, %r, %r, %r),\n" % (code, ar, en, dict(sorted(tp.items())))
        for code, ar, en, tp in rows
    )
    customer_lines = "".join("    (%r, %r),\n" % (tier, name) for tier, name in customers)
    return TEMPLATE.format(
        source=source,
        stamp=datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M:%S UTC"),
        tiers=list(BLOCKS),
        products=products,
        customers=customer_lines,
    )


def main(argv):
    if len(argv) != 2:
        sys.exit("usage: import_price_book.py <path to xlsx>")
    source = argv[1]
    customers, rows = build(source)
    target = __file__.replace("tools/import_price_book.py", "data/price_book.py")
    with open(target, "w", encoding="utf-8") as handle:
        handle.write(render(source, customers, rows))
    print("wrote %s: %d products, %d sheet customers" % (target, len(rows), len(customers)))


if __name__ == "__main__":
    main(sys.argv)
